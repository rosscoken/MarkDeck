"""
PowerPoint Presentation Generator

Generates PPTX files from parsed Markdown AST using python-pptx.
"""

import json
from pathlib import Path
from typing import Dict, Any, Optional, List
from dataclasses import dataclass

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


class MarkDeckError(Exception):
    """Base exception for MarkDeck errors."""
    pass


@dataclass
class SlideContent:
    """Represents content for a single slide."""
    title: str = ""
    content: List[Dict[str, Any]] = None
    speaker_notes: str = ""
    slide_type: str = "content"  # title, content, section
    
    def __post_init__(self):
        if self.content is None:
            self.content = []


class PPTXGenerator:
    """Generates PowerPoint presentations from Pandoc AST."""
    
    def __init__(self):
        self.theme_config = {}
        self.presentation = None
    
    def generate(
        self, 
        ast: Dict[str, Any], 
        output_path: Path,
        theme_config: Optional[Dict[str, Any]] = None,
        template_path: Optional[Path] = None
    ) -> None:
        """
        Generate PowerPoint presentation from AST.
        
        Args:
            ast: Pandoc AST dictionary
            output_path: Path for output PPTX file
            theme_config: Theme configuration dictionary
            template_path: Optional PowerPoint template file
        """
        self.theme_config = theme_config or self._get_default_theme()
        
        # Create presentation
        if template_path and template_path.exists():
            self.presentation = Presentation(str(template_path))
        else:
            self.presentation = Presentation()
        
        # Parse AST into slides
        slides = self._parse_ast_to_slides(ast)
        
        # Generate slides
        for slide_content in slides:
            self._create_slide(slide_content)
        
        # Save presentation
        self.presentation.save(str(output_path))
    
    def _get_default_theme(self) -> Dict[str, Any]:
        """Get default theme configuration."""
        return {
            "fonts": {
                "heading": {"family": "Calibri", "size": 32, "color": "#1f4e79"},
                "body": {"family": "Calibri", "size": 18, "color": "#333333"},
                "code": {"family": "Consolas", "size": 14, "color": "#000000"}
            },
            "colors": {
                "primary": "#1f4e79",
                "secondary": "#70ad47",
                "background": "#ffffff",
                "text": "#333333"
            },
            "layouts": {
                "title": 0,
                "title_content": 1,
                "section": 2
            }
        }
    
    def _parse_ast_to_slides(self, ast: Dict[str, Any]) -> List[SlideContent]:
        """Parse Pandoc AST into slide content."""
        slides = []
        current_slide = None
        blocks = ast.get("blocks", [])
        
        for block in blocks:
            block_type = block.get("t")
            block_content = block.get("c")
            
            if block_type == "Header":
                # New slide on header
                if current_slide:
                    slides.append(current_slide)
                
                level, attrs, inlines = block_content
                title = self._extract_text_from_inlines(inlines)
                
                # Determine slide type based on header level
                if level == 1:
                    slide_type = "section" if slides else "title"
                else:
                    slide_type = "content"
                
                current_slide = SlideContent(
                    title=title,
                    slide_type=slide_type
                )
            
            elif block_type == "Div":
                # Check for speaker notes
                attrs, content_blocks = block_content
                classes = attrs[1] if len(attrs) > 1 else []
                
                if "notes" in classes:
                    # Extract speaker notes
                    notes_text = self._extract_text_from_blocks(content_blocks)
                    if current_slide:
                        current_slide.speaker_notes = notes_text
                else:
                    # Regular div content
                    if current_slide:
                        current_slide.content.append(block)
            
            elif block_type == "HorizontalRule":
                # Slide break
                if current_slide:
                    slides.append(current_slide)
                    current_slide = SlideContent()
            
            else:
                # Regular content
                if current_slide is None:
                    current_slide = SlideContent()
                current_slide.content.append(block)
        
        # Add last slide if any
        if current_slide:
            slides.append(current_slide)
        
        return slides
    
    def _create_slide(self, slide_content: SlideContent) -> None:
        """Create a PowerPoint slide from slide content."""
        # Get layout based on slide type
        layout_idx = self._get_layout_index(slide_content.slide_type)
        slide_layout = self.presentation.slide_layouts[layout_idx]
        slide = self.presentation.slides.add_slide(slide_layout)
        
        # Set title
        if slide_content.title and slide.shapes.title:
            slide.shapes.title.text = slide_content.title
            self._apply_font_style(slide.shapes.title.text_frame, "heading")
        
        # Add content
        if slide_content.content:
            self._add_content_to_slide(slide, slide_content.content)
        
        # Add speaker notes
        if slide_content.speaker_notes:
            notes_slide = slide.notes_slide
            notes_slide.notes_text_frame.text = slide_content.speaker_notes
    
    def _get_layout_index(self, slide_type: str) -> int:
        """Get layout index for slide type."""
        layouts = self.theme_config.get("layouts", {})
        
        if slide_type == "title":
            return layouts.get("title", 0)
        elif slide_type == "section":
            return layouts.get("section", 2)
        else:
            return layouts.get("title_content", 1)
    
    def _add_content_to_slide(self, slide, content_blocks: List[Dict[str, Any]]) -> None:
        """Add content blocks to slide."""
        # Try to find content placeholder
        content_placeholder = None
        for shape in slide.placeholders:
            if hasattr(shape, 'text_frame') and shape.placeholder_format.idx == 1:
                content_placeholder = shape
                break
        
        if not content_placeholder:
            # Create a text box if no placeholder found
            left = Inches(1)
            top = Inches(2)
            width = Inches(8)
            height = Inches(5)
            content_placeholder = slide.shapes.add_textbox(left, top, width, height)
        
        # Process content blocks
        text_frame = content_placeholder.text_frame
        text_frame.clear()
        
        for block in content_blocks:
            self._add_block_to_textframe(text_frame, block)
    
    def _add_block_to_textframe(self, text_frame, block: Dict[str, Any]) -> None:
        """Add a content block to text frame."""
        block_type = block.get("t")
        block_content = block.get("c")
        
        if block_type == "Para":
            # Paragraph
            p = text_frame.paragraphs[0] if len(text_frame.paragraphs) == 1 and not text_frame.text else text_frame.add_paragraph()
            text = self._extract_text_from_inlines(block_content)
            p.text = text
            self._apply_font_style(text_frame, "body")
        
        elif block_type == "BulletList":
            # Bullet list
            for item in block_content:
                p = text_frame.add_paragraph()
                text = self._extract_text_from_blocks(item)
                p.text = text
                p.level = 0
                self._apply_font_style(text_frame, "body")
        
        elif block_type == "CodeBlock":
            # Code block
            attrs, code_text = block_content
            p = text_frame.add_paragraph()
            p.text = code_text
            self._apply_font_style(text_frame, "code")
        
        elif block_type == "Image":
            # Image - for MVP, just add alt text
            attrs, alt_text, target = block_content
            p = text_frame.add_paragraph()
            alt = self._extract_text_from_inlines(alt_text)
            p.text = f"[Image: {alt}]"
            self._apply_font_style(text_frame, "body")
    
    def _apply_font_style(self, text_frame, font_type: str) -> None:
        """Apply font styling from theme."""
        font_config = self.theme_config.get("fonts", {}).get(font_type, {})
        
        if not font_config:
            return
        
        for paragraph in text_frame.paragraphs:
            for run in paragraph.runs:
                if "family" in font_config:
                    run.font.name = font_config["family"]
                if "size" in font_config:
                    run.font.size = Pt(font_config["size"])
                if "color" in font_config:
                    color_hex = font_config["color"].lstrip("#")
                    rgb = tuple(int(color_hex[i:i+2], 16) for i in (0, 2, 4))
                    run.font.color.rgb = RGBColor(*rgb)
    
    def _extract_text_from_inlines(self, inlines: List[Dict[str, Any]]) -> str:
        """Extract plain text from Pandoc inline elements."""
        text_parts = []
        
        for inline in inlines:
            if not isinstance(inline, dict):
                continue
                
            inline_type = inline.get("t")
            inline_content = inline.get("c")
            
            if inline_type == "Str":
                text_parts.append(inline_content)
            elif inline_type == "Space":
                text_parts.append(" ")
            elif inline_type == "SoftBreak":
                text_parts.append(" ")
            elif inline_type == "LineBreak":
                text_parts.append("\n")
            elif inline_type in ["Emph", "Strong", "Strikeout"]:
                if isinstance(inline_content, list):
                    text_parts.append(self._extract_text_from_inlines(inline_content))
            elif inline_type == "Code":
                if isinstance(inline_content, list) and len(inline_content) >= 2:
                    text_parts.append(inline_content[1])
            elif inline_type == "Link":
                if isinstance(inline_content, list) and len(inline_content) >= 2:
                    text_parts.append(self._extract_text_from_inlines(inline_content[1]))
        
        return "".join(text_parts)
    
    def _extract_text_from_blocks(self, blocks: List[Dict[str, Any]]) -> str:
        """Extract plain text from Pandoc block elements."""
        text_parts = []
        
        for block in blocks:
            if not isinstance(block, dict):
                continue
                
            block_type = block.get("t")
            block_content = block.get("c")
            
            if block_type == "Para":
                if isinstance(block_content, list):
                    text_parts.append(self._extract_text_from_inlines(block_content))
            elif block_type == "Plain":
                if isinstance(block_content, list):
                    text_parts.append(self._extract_text_from_inlines(block_content))
        
        return " ".join(text_parts)