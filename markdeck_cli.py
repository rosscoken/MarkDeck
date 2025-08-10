#!/usr/bin/env python3
"""
MarkDeck CLI Standalone Entry Point

Standalone CLI that doesn't depend on complex module structure.
"""

import sys
from pathlib import Path
from typing import Optional
import shutil
import json
import subprocess
import logging
from dataclasses import dataclass

import typer
from rich.console import Console
import jsonschema

project_root = Path(__file__).parent

app = typer.Typer(
    name="markdeck",
    help="Convert Markdown files to PowerPoint presentations",
    add_completion=False,
)
console = Console()


class ParseError(Exception):
    """Exception raised when parsing fails."""
    pass


class MarkDeckError(Exception):
    """Base exception for MarkDeck errors."""
    pass


# Default theme content for init command
DEFAULT_THEME_JSON = """{
  "fonts": {
    "heading": {"family": "Calibri", "size": 32, "color": "#1f4e79"},
    "body": {"family": "Calibri", "size": 18, "color": "#333333"},
    "code": {"family": "Consolas", "size": 14, "color": "#000000"}
  },
  "colors": {
    "primary": "#1f4e79",
    "secondary": "#70ad47",
    "background": "#ffffff",
    "text": "#333333"
  }
}"""


class PandocBridge:
    """Bridge interface for Pandoc operations."""
    
    def __init__(self, pandoc_path: str = "pandoc"):
        self.pandoc_path = pandoc_path
        self._verify_pandoc()
    
    def _verify_pandoc(self) -> None:
        """Verify Pandoc is available and get version."""
        try:
            result = subprocess.run(
                [self.pandoc_path, "--version"],
                capture_output=True,
                text=True,
                check=True
            )
            version_line = result.stdout.split('\n')[0]
            logging.info(f"Found {version_line}")
        except (subprocess.CalledProcessError, FileNotFoundError):
            raise ParseError(f"Pandoc not found at '{self.pandoc_path}'. Please install Pandoc.")
    
    def convert_to_ast(self, content: str) -> dict:
        """Convert content to Pandoc AST."""
        try:
            result = subprocess.run(
                [self.pandoc_path, "--from", "markdown", "--to", "json"],
                input=content,
                capture_output=True,
                text=True,
                check=True
            )
            return json.loads(result.stdout)
        except subprocess.CalledProcessError as e:
            raise ParseError(f"Pandoc conversion failed: {e.stderr}")
        except json.JSONDecodeError as e:
            raise ParseError(f"Failed to parse Pandoc JSON output: {e}")


@dataclass
class SlideContent:
    """Represents content for a single slide."""
    title: str = ""
    content: list = None
    speaker_notes: str = ""
    slide_type: str = "content"
    
    def __post_init__(self):
        if self.content is None:
            self.content = []


class PPTXGenerator:
    """Generates PowerPoint presentations from Pandoc AST."""
    
    def __init__(self):
        # Import pptx here to avoid issues if not installed
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        
        self.Presentation = Presentation
        self.Inches = Inches
        self.Pt = Pt
        self.RGBColor = RGBColor
        self.theme_config = {}
        self.presentation = None
    
    def generate(self, ast: dict, output_path: Path, theme_config: Optional[dict] = None, template_path: Optional[Path] = None) -> None:
        """Generate PowerPoint presentation from AST."""
        self.theme_config = theme_config or self._get_default_theme()
        
        # Create presentation
        if template_path and template_path.exists():
            self.presentation = self.Presentation(str(template_path))
        else:
            self.presentation = self.Presentation()
        
        # Parse AST into slides
        slides = self._parse_ast_to_slides(ast)
        
        # Generate slides
        for slide_content in slides:
            self._create_slide(slide_content)
        
        # Save presentation
        self.presentation.save(str(output_path))
    
    def _get_default_theme(self) -> dict:
        """Get default theme configuration."""
        return {
            "fonts": {
                "heading": {"family": "Calibri", "size": 32, "color": "#1f4e79"},
                "body": {"family": "Calibri", "size": 18, "color": "#333333"},
                "code": {"family": "Consolas", "size": 14, "color": "#000000"}
            },
            "colors": {
                "primary": "#1f4e79",
                "secondary": "#70ad47",
                "background": "#ffffff",
                "text": "#333333"
            },
            "layouts": {
                "title": 0,
                "title_content": 1,
                "section": 2
            }
        }
    
    def _parse_ast_to_slides(self, ast: dict) -> list:
        """Parse Pandoc AST into slide content."""
        slides = []
        current_slide = None
        blocks = ast.get("blocks", [])
        
        for block in blocks:
            block_type = block.get("t")
            block_content = block.get("c")
            
            if block_type == "Header":
                # New slide on header
                if current_slide:
                    slides.append(current_slide)
                
                level, attrs, inlines = block_content
                title = self._extract_text_from_inlines(inlines)
                
                # Determine slide type based on header level
                if level == 1:
                    slide_type = "section" if slides else "title"
                else:
                    slide_type = "content"
                
                current_slide = SlideContent(
                    title=title,
                    slide_type=slide_type
                )
            
            elif block_type == "Div":
                # Check for speaker notes
                attrs, content_blocks = block_content
                classes = attrs[1] if len(attrs) > 1 else []
                
                if "notes" in classes:
                    # Extract speaker notes
                    notes_text = self._extract_text_from_blocks(content_blocks)
                    if current_slide:
                        current_slide.speaker_notes = notes_text
                else:
                    # Regular div content
                    if current_slide:
                        current_slide.content.append(block)
            
            elif block_type == "HorizontalRule":
                # Slide break
                if current_slide:
                    slides.append(current_slide)
                    current_slide = SlideContent()
            
            else:
                # Regular content
                if current_slide is None:
                    current_slide = SlideContent()
                current_slide.content.append(block)
        
        # Add last slide if any
        if current_slide:
            slides.append(current_slide)
        
        return slides
    
    def _create_slide(self, slide_content: SlideContent) -> None:
        """Create a PowerPoint slide from slide content."""
        # Get layout based on slide type
        layout_idx = self._get_layout_index(slide_content.slide_type)
        slide_layout = self.presentation.slide_layouts[layout_idx]
        slide = self.presentation.slides.add_slide(slide_layout)
        
        # Set title
        if slide_content.title and slide.shapes.title:
            slide.shapes.title.text = slide_content.title
            self._apply_font_style(slide.shapes.title.text_frame, "heading")
        
        # Add content
        if slide_content.content:
            self._add_content_to_slide(slide, slide_content.content)
        
        # Add speaker notes
        if slide_content.speaker_notes:
            notes_slide = slide.notes_slide
            notes_slide.notes_text_frame.text = slide_content.speaker_notes
    
    def _get_layout_index(self, slide_type: str) -> int:
        """Get layout index for slide type."""
        layouts = self.theme_config.get("layouts", {})
        
        if slide_type == "title":
            return layouts.get("title", 0)
        elif slide_type == "section":
            return layouts.get("section", 2)
        else:
            return layouts.get("title_content", 1)
    
    def _add_content_to_slide(self, slide, content_blocks: list) -> None:
        """Add content blocks to slide."""
        # Try to find content placeholder
        content_placeholder = None
        for shape in slide.placeholders:
            if hasattr(shape, 'text_frame') and shape.placeholder_format.idx == 1:
                content_placeholder = shape
                break
        
        if not content_placeholder:
            # Create a text box if no placeholder found
            left = self.Inches(1)
            top = self.Inches(2)
            width = self.Inches(8)
            height = self.Inches(5)
            content_placeholder = slide.shapes.add_textbox(left, top, width, height)
        
        # Process content blocks
        text_frame = content_placeholder.text_frame
        text_frame.clear()
        
        for block in content_blocks:
            self._add_block_to_textframe(text_frame, block)
    
    def _add_block_to_textframe(self, text_frame, block: dict) -> None:
        """Add a content block to text frame."""
        block_type = block.get("t")
        block_content = block.get("c")
        
        if block_type == "Para":
            # Paragraph
            p = text_frame.paragraphs[0] if len(text_frame.paragraphs) == 1 and not text_frame.text else text_frame.add_paragraph()
            text = self._extract_text_from_inlines(block_content)
            p.text = text
            self._apply_font_style(text_frame, "body")
        
        elif block_type == "BulletList":
            # Bullet list
            for item in block_content:
                p = text_frame.add_paragraph()
                text = self._extract_text_from_blocks(item)
                p.text = text
                p.level = 0
                self._apply_font_style(text_frame, "body")
        
        elif block_type == "CodeBlock":
            # Code block
            attrs, code_text = block_content
            p = text_frame.add_paragraph()
            p.text = code_text
            self._apply_font_style(text_frame, "code")
        
        elif block_type == "Image":
            # Image - for MVP, just add alt text
            attrs, alt_text, target = block_content
            p = text_frame.add_paragraph()
            alt = self._extract_text_from_inlines(alt_text)
            p.text = f"[Image: {alt}]"
            self._apply_font_style(text_frame, "body")
    
    def _apply_font_style(self, text_frame, font_type: str) -> None:
        """Apply font styling from theme."""
        font_config = self.theme_config.get("fonts", {}).get(font_type, {})
        
        if not font_config:
            return
        
        for paragraph in text_frame.paragraphs:
            for run in paragraph.runs:
                if "family" in font_config:
                    run.font.name = font_config["family"]
                if "size" in font_config:
                    run.font.size = self.Pt(font_config["size"])
                if "color" in font_config:
                    color_hex = font_config["color"].lstrip("#")
                    rgb = tuple(int(color_hex[i:i+2], 16) for i in (0, 2, 4))
                    run.font.color.rgb = self.RGBColor(*rgb)
    
    def _extract_text_from_inlines(self, inlines: list) -> str:
        """Extract plain text from Pandoc inline elements."""
        text_parts = []
        
        for inline in inlines:
            if not isinstance(inline, dict):
                continue
                
            inline_type = inline.get("t")
            inline_content = inline.get("c")
            
            if inline_type == "Str":
                text_parts.append(inline_content)
            elif inline_type == "Space":
                text_parts.append(" ")
            elif inline_type == "SoftBreak":
                text_parts.append(" ")
            elif inline_type == "LineBreak":
                text_parts.append("\n")
            elif inline_type in ["Emph", "Strong", "Strikeout"]:
                if isinstance(inline_content, list):
                    text_parts.append(self._extract_text_from_inlines(inline_content))
            elif inline_type == "Code":
                if isinstance(inline_content, list) and len(inline_content) >= 2:
                    text_parts.append(inline_content[1])
            elif inline_type == "Link":
                if isinstance(inline_content, list) and len(inline_content) >= 2:
                    text_parts.append(self._extract_text_from_inlines(inline_content[1]))
        
        return "".join(text_parts)
    
    def _extract_text_from_blocks(self, blocks: list) -> str:
        """Extract plain text from Pandoc block elements."""
        text_parts = []
        
        for block in blocks:
            if not isinstance(block, dict):
                continue
                
            block_type = block.get("t")
            block_content = block.get("c")
            
            if block_type == "Para":
                if isinstance(block_content, list):
                    text_parts.append(self._extract_text_from_inlines(block_content))
            elif block_type == "Plain":
                if isinstance(block_content, list):
                    text_parts.append(self._extract_text_from_inlines(block_content))
        
        return " ".join(text_parts)


def detect_pandoc() -> str:
    """Detect if pandoc is installed and return path."""
    pandoc_path = shutil.which("pandoc")
    if not pandoc_path:
        console.print("[red]Error: Pandoc not found![/red]")
        console.print("\nPandoc is required for MarkDeck to work.")
        console.print("Please install Pandoc:")
        console.print("• Ubuntu/Debian: [bold]sudo apt install pandoc[/bold]")
        console.print("• macOS: [bold]brew install pandoc[/bold]")
        console.print("• Windows: [bold]choco install pandoc[/bold]")
        console.print("• Or download from: https://pandoc.org/installing.html")
        raise typer.Exit(1)
    return pandoc_path


def validate_theme_file(theme_path: Path) -> bool:
    """Validate a theme file against the schema."""
    try:
        # Load schema
        schema_path = project_root / "schemas" / "theme.schema.json"
        if not schema_path.exists():
            console.print(f"[red]Error: Schema not found at {schema_path}[/red]")
            return False
            
        with open(schema_path) as f:
            schema = json.load(f)
        
        # Load theme
        with open(theme_path) as f:
            theme = json.load(f)
        
        # Validate
        jsonschema.validate(theme, schema)
        return True
        
    except FileNotFoundError:
        console.print(f"[red]Error: Theme file not found: {theme_path}[/red]")
        return False
    except json.JSONDecodeError as e:
        console.print(f"[red]Error: Invalid JSON in theme file: {e}[/red]")
        return False
    except jsonschema.ValidationError as e:
        console.print(f"[red]Error: Theme validation failed: {e.message}[/red]")
        if e.absolute_path:
            console.print(f"[red]Path: {'.'.join(map(str, e.absolute_path))}[/red]")
        return False
    except Exception as e:
        console.print(f"[red]Error: {e}[/red]")
        return False


@app.command()
def build(
    input_file: Path = typer.Argument(..., help="Input Markdown file"),
    output: Optional[Path] = typer.Option(None, "-o", "--output", help="Output PowerPoint file"),
    theme: Optional[Path] = typer.Option(None, "--theme", help="Theme JSON file"),
    template: Optional[Path] = typer.Option(None, "--template", help="PowerPoint template file"),
    config: Optional[Path] = typer.Option(None, "--config", help="Config file (markdeck.toml/yaml)"),
    verbose: bool = typer.Option(False, "--verbose", "-v", help="Verbose output"),
):
    """Build a PowerPoint presentation from Markdown."""
    
    if verbose:
        console.print(f"[blue]Building presentation from {input_file}[/blue]")
    
    # Check if input file exists
    if not input_file.exists():
        console.print(f"[red]Error: Input file not found: {input_file}[/red]")
        raise typer.Exit(1)
    
    # Determine output path
    if output is None:
        output = input_file.with_suffix(".pptx")
    
    # Detect pandoc
    pandoc_path = detect_pandoc()
    
    # Validate theme if provided
    if theme and not validate_theme_file(theme):
        raise typer.Exit(1)
    
    try:
        # Initialize components
        bridge = PandocBridge(pandoc_path)
        generator = PPTXGenerator()
        
        # Read and parse markdown
        content = input_file.read_text(encoding='utf-8')
        if verbose:
            console.print("[blue]Parsing Markdown with Pandoc...[/blue]")
        
        ast = bridge.convert_to_ast(content)
        
        # Load theme if provided
        theme_config = None
        if theme:
            with open(theme) as f:
                theme_config = json.load(f)
        
        # Generate presentation
        if verbose:
            console.print("[blue]Generating PowerPoint presentation...[/blue]")
        
        generator.generate(ast, output, theme_config, template)
        
        console.print(f"[green]✓ Generated presentation: {output}[/green]")
        
    except Exception as e:
        console.print(f"[red]Error: {e}[/red]")
        if verbose:
            import traceback
            console.print(f"[red]{traceback.format_exc()}[/red]")
        raise typer.Exit(1)


@app.command()
def init(
    force: bool = typer.Option(False, "--force", help="Overwrite existing files"),
    directory: Path = typer.Option(Path("."), "--dir", help="Directory to create files in"),
):
    """Initialize example files in the current directory."""
    
    # Create examples directory if it doesn't exist
    examples_dir = directory / "examples"
    examples_dir.mkdir(exist_ok=True)
    
    # Paths for example files
    basic_md = examples_dir / "basic.md"
    theme_json = examples_dir / "theme.json"
    
    # Check for existing files
    if basic_md.exists() and not force:
        console.print(f"[yellow]Warning: {basic_md} already exists. Use --force to overwrite.[/yellow]")
    if theme_json.exists() and not force:
        console.print(f"[yellow]Warning: {theme_json} already exists. Use --force to overwrite.[/yellow]")
    
    if (basic_md.exists() or theme_json.exists()) and not force:
        console.print("[yellow]Use --force to overwrite existing files.[/yellow]")
        raise typer.Exit(1)
    
    try:
        # Get source files from package
        source_dir = project_root / "examples"
        
        # Copy or create files
        if source_dir.exists():
            source_basic = source_dir / "basic.md"
            source_theme = source_dir / "theme.json"
            
            if source_basic.exists():
                shutil.copy2(source_basic, basic_md)
                console.print(f"[green]✓ Created {basic_md}[/green]")
            
            if source_theme.exists():
                shutil.copy2(source_theme, theme_json)
                console.print(f"[green]✓ Created {theme_json}[/green]")
        
        # Create files if they don't exist yet
        if not basic_md.exists():
            basic_content = """# Welcome to MarkDeck

## Getting Started

This is a basic MarkDeck presentation.

- Create content in Markdown
- Build with `markdeck build`
- Customize with themes

::: notes
This is a speaker note that will be included in the presenter notes.
:::

## Next Steps

Run `markdeck build examples/basic.md` to generate your first presentation!
"""
            basic_md.write_text(basic_content)
            console.print(f"[green]✓ Created {basic_md}[/green]")
        
        if not theme_json.exists():
            theme_json.write_text(DEFAULT_THEME_JSON)
            console.print(f"[green]✓ Created {theme_json}[/green]")
            
        console.print("\n[blue]Example files created! Try:[/blue]")
        console.print(f"[bold]markdeck build {basic_md} --theme {theme_json}[/bold]")
        
    except Exception as e:
        console.print(f"[red]Error creating example files: {e}[/red]")
        raise typer.Exit(1)


@app.command()
def validate(
    theme: Path = typer.Option(..., "--theme", help="Theme JSON file to validate"),
):
    """Validate a theme configuration file."""
    
    console.print(f"[blue]Validating theme: {theme}[/blue]")
    
    if validate_theme_file(theme):
        console.print(f"[green]✓ Theme {theme} is valid![/green]")
    else:
        raise typer.Exit(1)


def version_callback(value: bool):
    """Handle version option."""
    if value:
        console.print("MarkDeck 0.1.0")
        raise typer.Exit()


@app.callback()
def main(
    version: Optional[bool] = typer.Option(None, "--version", callback=version_callback, is_eager=True, help="Show version"),
):
    """MarkDeck: Convert Markdown to PowerPoint presentations."""
    pass


if __name__ == "__main__":
    app()